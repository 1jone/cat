from pathlib import Path
import sys
from docx import Document


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: extract_docx_text.py <docx>")
    path = Path(sys.argv[1])
    doc = Document(path)
    print(f"FILE: {path}")
    print("=== PARAGRAPHS ===")
    for i, p in enumerate(doc.paragraphs, 1):
        text = p.text.strip()
        if text:
            print(f"P{i}: {text}")
    print("=== TABLES ===")
    for ti, table in enumerate(doc.tables, 1):
        print(f"[TABLE {ti}]")
        for ri, row in enumerate(table.rows, 1):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                print(f"R{ri}: " + " | ".join(cells))


if __name__ == "__main__":
    main()

from pathlib import Path
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output"
OUT_DIR.mkdir(exist_ok=True)
ORIGINAL = Path(r"D:\zhuomian\王芬艳简历.docx")
PHOTO_DIR = OUT_DIR / "resume_media_extract"
PHOTO_DIR.mkdir(exist_ok=True)
OUT_FILE = OUT_DIR / "王芬艳_前端开发工程师_简历_一页优化版.docx"

BLUE = "1F4E79"
LIGHT_BLUE = "EAF2FB"
DARK = "1A1A1A"
GRAY = "555555"
MID_GRAY = "D9E2F3"


def extract_photo():
    if not ORIGINAL.exists():
        return None
    candidates = []
    with zipfile.ZipFile(ORIGINAL) as zf:
        for name in zf.namelist():
            if name.startswith("word/media/"):
                data = zf.read(name)
                suffix = Path(name).suffix.lower() or ".png"
                out = PHOTO_DIR / f"resume_photo_{len(candidates)+1}{suffix}"
                out.write_bytes(data)
                candidates.append((out, len(data)))
    if not candidates:
        return None
    usable = [c for c in candidates if c[0].suffix.lower() in (".jpg", ".jpeg", ".png") and c[1] > 1024]
    if not usable:
        return None
    source = max(usable, key=lambda x: x[1])[0]
    normalized = PHOTO_DIR / "resume_photo_normalized.png"
    with Image.open(source) as im:
        im.convert("RGB").save(normalized, "PNG")
    return normalized


def set_run_font(run, size=10, color=DARK, bold=None):
    name = "Microsoft YaHei"
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def remove_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)


def mark_row_as_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.05)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.25)
    section.right_margin = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.02

    for style_name in ("List Bullet", "List Bullet 2"):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(8.5)
        style.paragraph_format.space_after = Pt(0.8)
        style.paragraph_format.line_spacing = 0.98


def para(cell_or_doc, text="", size=9.2, color=DARK, bold=False, after=2, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell_or_doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.02
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)
    return p


def heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run_font(r, size=10.6, color=BLUE, bold=True)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), MID_GRAY)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.42)
    p.paragraph_format.first_line_indent = Cm(-0.18)
    p.paragraph_format.space_after = Pt(0.8)
    r = p.add_run(text)
    set_run_font(r, size=8.5)


def project(doc, name, stack, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0.8)
    r = p.add_run(name)
    set_run_font(r, size=9.4, bold=True)
    r2 = p.add_run(f"  |  {stack}")
    set_run_font(r2, size=8.2, color=GRAY)
    for item in bullets:
        bullet(doc, item)


def build():
    photo = extract_photo()
    doc = Document()
    configure(doc)

    header = doc.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Cm(14.2)
    header.columns[1].width = Cm(3.8)
    mark_row_as_header(header.rows[0])
    left, right = header.rows[0].cells
    for cell in (left, right):
        remove_cell_borders(cell)
        set_cell_margins(cell, top=60, bottom=60, start=90, end=90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(left, LIGHT_BLUE)
    set_cell_shading(right, LIGHT_BLUE)

    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("王芬艳")
    set_run_font(r, size=18, color=BLUE, bold=True)
    para(left, "前端开发工程师 / 小程序开发工程师", size=10.5, color=DARK, bold=True, after=3)
    para(left, "手机：17755787620  |  学历：专科  |  性别：女  |  出生年月：2004.05  |  政治面貌：团员", size=8.8, color=GRAY, after=1)
    para(left, "关键词：Vue 3 / TypeScript / uni-app / 微信小程序 / Ant Design Vue / ECharts / Vite", size=8.8, color=GRAY, after=0)

    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if photo and photo.exists():
        rr = rp.add_run()
        inline = rr.add_picture(str(photo), width=Cm(2.85))
        inline._inline.docPr.set("descr", "王芬艳证件照")
        inline._inline.docPr.set("title", "王芬艳证件照")
    else:
        rr = rp.add_run("照片")
        set_run_font(rr, size=10, color=GRAY, bold=True)

    heading(doc, "个人优势")
    for item in [
        "熟悉 Vue 3、TypeScript、uni-app、小程序与后台管理系统开发，能独立完成页面、组件、接口联调和发布。",
        "在小程序商城、企业官网/后台、抖音小游戏项目中承担项目推进、需求沟通、测试验收与交付协调工作。",
        "有竞赛经历和多项目落地经验，关注用户体验、性能优化、问题闭环和代码可维护性。",
    ]:
        bullet(doc, item)

    heading(doc, "专业技能")
    for item in [
        "前端：Vue 3、TypeScript、Nuxt.js、JavaScript ES6+、HTML5、CSS3、响应式布局。",
        "跨端/小程序：uni-app、微信小程序原生开发、授权、支付、多端适配与发布上架。",
        "工程化/UI：Vite、Webpack、Ant Design Vue、Element UI、TailwindCSS、Bootstrap。",
        "数据与通信：ECharts、WebSocket、RESTful API、数据大屏、自适应布局。",
    ]:
        bullet(doc, item)

    heading(doc, "工作经历")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("安徽易搜有限科技公司")
    set_run_font(r, size=9.8, bold=True)
    r2 = p.add_run("  |  前端开发 / 项目推进 / 测试验收  |  2025.05.05-至今")
    set_run_font(r2, size=8.4, color=GRAY)
    bullet(doc, "负责公司 Web、小程序、后台管理系统及小游戏相关项目的前端开发，同时参与需求沟通、任务拆分、测试回归和交付上线。")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("代表项目：")
    set_run_font(r, size=9.2, color=BLUE, bold=True)
    project(
        doc,
        "微信小程序商城及管理系统",
        "uni-app + Vue 3 + TypeScript + Ant Design Vue",
        [
            "担任项目推进与核心开发，梳理商品、订单、库存、营销活动、后台管理等功能模块，拆分开发任务并跟进进度。",
            "完成小程序端与管理后台页面开发、接口联调、微信授权/支付接入及多端适配，支持小程序/H5/App 复用。",
            "组织功能测试与问题回归，跟进订单状态、库存同步、支付流程等关键链路，协助完成发布上线。",
        ],
    )
    project(
        doc,
        "数据可视化大屏系统",
        "Vue 3 + ECharts + WebSocket",
        [
            "通过 WebSocket 实现实时数据推送，使用 ECharts 开发联动图表，并适配多种大屏分辨率。",
        ],
    )
    project(
        doc,
        "扬子冷库官网及管理系统",
        "Vue 3 + TypeScript + Ant Design Vue",
        [
            "负责项目需求沟通、页面规划和开发排期，推进企业官网与后台管理系统从开发、联调到验收交付。",
            "完成官网响应式展示、后台权限管理、内容发布、数据统计等模块开发，与后端 RESTful API 联调。",
            "参与测试用例整理、兼容性检查和问题修复，保障后台操作流程、权限控制和页面展示稳定可用。",
        ],
    )
    project(
        doc,
        "猫咪追追追 - 抖音小游戏",
        "原生 JavaScript + Canvas + tt.* API",
        [
            "参与项目规划、功能拆分、测试验收和版本发布，协调玩法、广告、侧边栏、排行榜等功能模块落地。",
            "完成 Canvas 目标动画、点击反馈、资源加载、状态切换及部分抖音平台能力接入，配合测试修复交互与兼容问题。",
        ],
    )

    heading(doc, "竞赛与证书")
    for item in [
        "2024.04 中国大学生计算机设计大赛安徽省级一等奖；2024.08 “中国软件杯”大学生软件设计大赛三等奖。",
        "2024.01、2025.01 安徽省职业院校技能大赛移动应用设计与开发二等奖；2023.04 中国大学生计算机设计大赛安徽省级三等奖。",
        "WPS 初级证书、1+X 初级/中级证书、计算机一级、计算机程序三级、英语 A 级。",
    ]:
        bullet(doc, item)

    doc.core_properties.title = "王芬艳 前端开发工程师 简历"
    doc.core_properties.author = "王芬艳"
    doc.core_properties.keywords = "前端开发, Vue 3, TypeScript, uni-app, 微信小程序, Ant Design Vue, ECharts"
    doc.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    print(build())

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output"
OUT_DIR.mkdir(exist_ok=True)
OUT_FILE = OUT_DIR / "王芬艳_前端开发工程师_简历_优化版.docx"

BLUE = "1F4E79"
DARK = "111111"
GRAY = "555555"


def set_run_font(run, name="Microsoft YaHei", size=10.5, color=DARK, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_bottom_border(paragraph, color=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_doc(doc):
    section = doc.sections[0]
    # Named override for Chinese resume delivery: A4 page instead of US Letter.
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    for style_name in ("List Bullet", "List Bullet 2"):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.2)
        style.paragraph_format.space_after = Pt(2.5)
        style.paragraph_format.line_spacing = 1.08

    for i in (1, 2, 3):
        style = doc.styles[f"Heading {i}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.font.bold = True
        style.font.size = Pt(12.5 if i == 1 else 11.5)
        style.paragraph_format.space_before = Pt(8 if i == 1 else 5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True


def add_para(doc, text="", size=10.5, color=DARK, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)
    return p


def add_section(doc, title):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(title)
    for run in p.runs:
        set_run_font(run, size=12.5, color=BLUE, bold=True)
    add_bottom_border(p)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.2)
    p.paragraph_format.space_after = Pt(2.5)
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def add_project(doc, title, meta, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1.5)
    r = p.add_run(title)
    set_run_font(r, size=10.8, color=DARK, bold=True)
    if meta:
        r2 = p.add_run("  |  " + meta)
        set_run_font(r2, size=9.5, color=GRAY)
    for bullet in bullets:
        add_bullet(doc, bullet)


def build():
    doc = Document()
    configure_doc(doc)

    p = add_para(doc, "王芬艳", size=20, color=BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    p.paragraph_format.space_before = Pt(0)
    add_para(
        doc,
        "求职意向：前端开发工程师 / 小程序开发工程师  |  手机：17755787620  |  学历：专科",
        size=9.8,
        color=GRAY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
    )
    add_bottom_border(doc.paragraphs[-1], color="D9E2F3")

    add_section(doc, "个人优势")
    for item in [
        "熟悉 Vue 3、TypeScript、Nuxt.js、uni-app、微信小程序与原生 JavaScript 开发，具备 Web、H5、小程序及跨端项目经验。",
        "能够独立完成前端页面搭建、组件封装、接口联调、数据可视化、权限管理、移动端适配与上线发布等工作。",
        "有 Canvas/小游戏项目实践，熟悉抖音小游戏 API、资源加载、游戏状态管理、广告/侧边栏/排行榜等平台能力接入。",
        "具备良好的沟通协作能力，能将业务需求拆解为可交付的前端功能，并关注性能、可维护性和用户体验。",
    ]:
        add_bullet(doc, item)

    add_section(doc, "专业技能")
    skill_lines = [
        "前端框架：Vue 3、TypeScript、Nuxt.js（SSR/SSG）、JavaScript ES6+、HTML5、CSS3。",
        "小程序/跨端：uni-app、微信小程序原生开发、抖音小游戏 tt.* API、授权、支付、发布上架与多端适配。",
        "UI 与样式：Ant Design Vue、Element UI、TailwindCSS、Bootstrap、响应式布局、移动端适配。",
        "工程化：Vite、Webpack、模块化开发、组件封装、前后端分离、RESTful API 联调、代码规范与性能优化。",
        "可视化与实时通信：ECharts、WebSocket、数据大屏自适应布局与低延迟数据更新。",
        "工具与证书：WPS 初级证书、1+X 初级/中级证书、计算机一级、计算机程序三级、英语 A 级。",
    ]
    for item in skill_lines:
        add_bullet(doc, item)

    add_section(doc, "项目经历")
    add_project(
        doc,
        "猫咪追追追 - 抖音小游戏",
        "原生 JavaScript + Canvas + 抖音小游戏 API（tt.*）",
        [
            "参与一款面向养猫家庭的屏幕互动小游戏开发，采用 Canvas + 原生 JavaScript 实现目标绘制、点击交互、动画更新和游戏主循环。",
            "负责/参与分层管理器架构梳理与功能实现，包含 Game 主协调器、InputManager、ResourceManager、GameStateManager、SpawnManager、AudioManager 等模块。",
            "实现多种目标运动与猫咪互动玩法，支持弹跳、随机、波浪、悬停、圆周、螺旋、8 字形等运动模式，并处理受惊逃离、速度恢复和边界同步等交互细节。",
            "接入抖音小游戏平台能力，包括广告激励、Banner/插屏频控、侧边栏奖励、好友排行榜、签到、金币/体力体系与本地存储，增强留存和商业化能力。",
            "围绕猫咪实际使用场景优化体验，强调 PLAYING 状态纯净互动区，避免猫掌误触广告或 UI，并补充设备固定、短时使用等安全运营方案。",
        ],
    )
    add_project(
        doc,
        "微信小程序商城及管理系统",
        "uni-app + Vue 3 + TypeScript + Ant Design Vue",
        [
            "负责微信小程序商城及配套管理后台开发，支持商品购买、订单管理、营销活动、商品管理与订单处理等核心功能。",
            "采用 uni-app 实现小程序、H5 与 App 多端兼容，提升代码复用率，并完成微信小程序平台发布上架。",
            "集成微信支付、用户授权等原生能力，配合后台接口实现库存更新与订单状态同步，提升多端协同效率。",
        ],
    )
    add_project(
        doc,
        "数据可视化大屏系统",
        "Vue 3 + ECharts + WebSocket",
        [
            "负责实时数据监控大屏前端开发，通过 WebSocket 长连接实现多源数据动态推送与低延迟更新。",
            "使用 ECharts 定制可钻取、联动的复杂图表，并设计自适应布局方案，适配多种大屏分辨率展示。",
        ],
    )
    add_project(
        doc,
        "扬子冷库官网及管理系统",
        "Vue 3 + TypeScript + Ant Design Vue",
        [
            "负责企业官网与后台管理系统开发，官网采用响应式设计实现多端适配，完整展示企业信息与产品内容。",
            "后台基于 Ant Design Vue 构建用户权限管理、内容发布、数据统计等模块，通过 RESTful API 与后端协作提升可维护性。",
        ],
    )

    add_section(doc, "竞赛与获奖")
    for item in [
        "2025.01 安徽省职业院校技能大赛移动应用设计与开发二等奖",
        "2024.08 “中国软件杯”大学生软件设计大赛三等奖",
        "2024.04 中国大学生计算机设计大赛安徽省级一等奖",
        "2024.01 安徽省职业院校技能大赛移动应用设计与开发二等奖",
        "2023.04 中国大学生计算机设计大赛安徽省级三等奖",
    ]:
        add_bullet(doc, item)

    add_section(doc, "基本信息")
    add_para(doc, "性别：女  |  出生年月：2004.05  |  政治面貌：团员", size=10.2, color=DARK, after=2)

    doc.core_properties.title = "王芬艳 前端开发工程师 简历"
    doc.core_properties.subject = "前端开发、Vue、TypeScript、小程序、抖音小游戏、Canvas"
    doc.core_properties.author = "王芬艳"
    doc.core_properties.keywords = "前端开发, Vue 3, TypeScript, uni-app, 微信小程序, 抖音小游戏, Canvas, ECharts"
    doc.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    print(build())
